"""Phase 8: admin per-student reports (ARCHITECTURE.md §5.3) -- the list view is computed with
set-based SQL aggregates over the current page of students (never one query per student, which is
what the documented prototype mistake looked like); the detail view is naturally one student.
"""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models.courses import Course, Lesson, LessonProgress, Section
from app.models.exams import AttemptStatus, Exam, ExamAttempt
from app.models.groups import Enrollment, GroupCourseEnrollment, GroupMembership
from app.models.identity import StudentProfile, User
from app.models.points import PointLedger
from app.services import live_sessions as live_svc
from app.services import points as points_svc
from app.services.access import accessible_course_ids, effective_subscription_status
from app.services.excel_style import write_report_sheet
from app.services.lessons import course_progress_percent
from app.services.students import get_student_or_404, student_filter_stmt
from app.schemas.common import Page
from app.schemas.points import PointEvent
from app.schemas.reports import (
    CourseProgressOut,
    ExamSummaryOut,
    PeriodExamOut,
    PeriodLessonOut,
    PeriodOut,
    ReportAttendanceOut,
    StudentPeriodReport,
    StudentReportDetail,
    StudentReportRow,
    StudentReportStudentOut,
)
from app.schemas.students import StudentGroupSummary


def _report_student_out(user: User, profile: StudentProfile) -> StudentReportStudentOut:
    return StudentReportStudentOut(
        id=user.id, full_name=user.full_name, student_code=user.student_code,
        grade_level=profile.grade_level.value if profile.grade_level else None,
        student_type=profile.student_type.value, effective_subscription=effective_subscription_status(profile),
    )


def list_student_reports(
    db: Session, academy_id: int, *, q: str | None, grade: str | None, student_type: str | None,
    subscription: str | None, group_id: int | None, course_id: int | None, sort: str | None, page: int, page_size: int,
) -> Page[StudentReportRow]:
    stmt = student_filter_stmt(academy_id, q=q, grade=grade, student_type=student_type, subscription=subscription, group_id=group_id, course_id=course_id)
    total = db.scalar(select(func.count()).select_from(stmt.with_only_columns(User.id).subquery())) or 0

    sort_map = {
        "name": User.full_name.asc(), "-name": User.full_name.desc(),
        "points": StudentProfile.total_points.desc(), "-points": StudentProfile.total_points.asc(),
    }
    stmt = stmt.order_by(sort_map.get(sort or "name", User.full_name.asc())).offset((page - 1) * page_size).limit(page_size)
    rows = db.execute(stmt).all()
    if not rows:
        return Page[StudentReportRow](items=[], total=total, page=page, page_size=page_size)

    users = {r.User.id: r.User for r in rows}
    profiles = {r.User.id: r.StudentProfile for r in rows}
    user_ids = list(users.keys())

    # -- accessible course ids per student (direct + group), for lesson-progress totals below --
    direct_pairs = db.execute(select(Enrollment.user_id, Enrollment.course_id).where(Enrollment.user_id.in_(user_ids))).all()
    group_pairs = db.execute(
        select(GroupMembership.user_id, GroupCourseEnrollment.course_id)
        .join(GroupCourseEnrollment, GroupCourseEnrollment.group_id == GroupMembership.group_id)
        .where(GroupMembership.user_id.in_(user_ids))
    ).all()
    course_ids_by_user: dict[int, set[int]] = {}
    for uid, cid in [*direct_pairs, *group_pairs]:
        course_ids_by_user.setdefault(uid, set()).add(cid)
    all_course_ids = {cid for cids in course_ids_by_user.values() for cid in cids}

    lesson_counts_by_course: dict[int, int] = {}
    if all_course_ids:
        lesson_counts_by_course = dict(db.execute(
            select(Section.course_id, func.count(Lesson.id)).select_from(Section)
            .join(Lesson, Lesson.section_id == Section.id).where(Section.course_id.in_(all_course_ids)).group_by(Section.course_id)
        ).all())
    total_lessons_by_user = {uid: sum(lesson_counts_by_course.get(cid, 0) for cid in cids) for uid, cids in course_ids_by_user.items()}

    completed_by_user: dict[int, int] = {}
    if all_course_ids:
        completed_by_user = dict(db.execute(
            select(LessonProgress.user_id, func.count(LessonProgress.id)).select_from(LessonProgress)
            .join(Lesson, Lesson.id == LessonProgress.lesson_id).join(Section, Section.id == Lesson.section_id)
            .where(LessonProgress.user_id.in_(user_ids), LessonProgress.completed.is_(True), Section.course_id.in_(all_course_ids))
            .group_by(LessonProgress.user_id)
        ).all())

    exam_rows = db.execute(
        select(ExamAttempt.user_id, func.count(func.distinct(ExamAttempt.exam_id)), func.avg(ExamAttempt.percentage))
        .where(ExamAttempt.user_id.in_(user_ids), ExamAttempt.status.in_((AttemptStatus.submitted, AttemptStatus.expired)))
        .group_by(ExamAttempt.user_id)
    ).all()
    exams_taken_by_user = {row[0]: row[1] for row in exam_rows}
    exam_avg_by_user = {row[0]: float(row[2]) for row in exam_rows if row[2] is not None}

    attendance_by_user = live_svc.attendance_summary_by_user(db, user_ids)

    items = []
    for r in rows:
        uid = r.User.id
        total_lessons = total_lessons_by_user.get(uid, 0)
        completed_lessons = completed_by_user.get(uid, 0)
        overall_progress = round((completed_lessons / total_lessons) * 100) if total_lessons else 0
        att = attendance_by_user.get(uid, {"present": 0, "late": 0, "absent": 0, "excused": 0, "rate": 0})
        items.append(StudentReportRow(
            student=_report_student_out(r.User, r.StudentProfile),
            overall_progress=overall_progress, courses_count=len(course_ids_by_user.get(uid, set())),
            completed_lessons=completed_lessons, total_lessons=total_lessons,
            exams_taken=exams_taken_by_user.get(uid, 0), exam_average=round(exam_avg_by_user[uid], 1) if uid in exam_avg_by_user else None,
            attendance=ReportAttendanceOut(present=att["present"], late=att["late"], absent=att["absent"], excused=att["excused"], rate=att["rate"]),
            total_points=r.StudentProfile.total_points or 0, last_login_at=r.User.last_login_at,
        ))
    return Page[StudentReportRow](items=items, total=total, page=page, page_size=page_size)


def export_student_reports_xlsx(rows: list[StudentReportRow], *, academy_name: str = "Talent Academy") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Students Report"
    headers = ["Student ID", "Full Name", "Grade", "Type", "Subscription", "Progress %", "Courses", "Lessons", "Exams Taken", "Exam Avg %", "Attendance %", "Points", "Last Login"]
    data = [
        [
            row.student.student_code or "", row.student.full_name, row.student.grade_level or "", row.student.student_type,
            row.student.effective_subscription, row.overall_progress, row.courses_count, f"{row.completed_lessons}/{row.total_lessons}",
            row.exams_taken, row.exam_average if row.exam_average is not None else "", row.attendance.rate, row.total_points,
            row.last_login_at.strftime("%Y-%m-%d %H:%M") if row.last_login_at else "",
        ]
        for row in rows
    ]
    totals = None
    if rows:
        avg_progress = round(sum(r.overall_progress for r in rows) / len(rows))
        avg_attendance = round(sum(r.attendance.rate for r in rows) / len(rows))
        total_points = sum(r.total_points for r in rows)
        totals = ["", f"Total: {len(rows)} students", "", "", "", avg_progress, "", "", "", "", avg_attendance, total_points, ""]
    write_report_sheet(
        ws, report_title="Students Report", headers=headers, rows=data, academy_name=academy_name, totals_row=totals,
        column_widths=[14, 22, 8, 10, 12, 11, 9, 12, 12, 11, 13, 9, 16],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------------- detail ----

def student_report_detail(db: Session, academy_id: int, student_id: int) -> StudentReportDetail:
    student = get_student_or_404(db, academy_id, student_id)
    profile = student.student_profile

    groups = [StudentGroupSummary(id=m.group.id, name=m.group.name) for m in sorted(student.group_memberships, key=lambda m: m.group.name.lower())]

    course_ids = accessible_course_ids(db, student.id)
    titles = dict(db.execute(select(Course.id, Course.title).where(Course.id.in_(course_ids))).all()) if course_ids else {}
    courses = []
    for course_id in course_ids:
        progress = course_progress_percent(db, student.id, course_id)
        rank_result = points_svc.course_rank_for(db, course_id, student.id)
        rank, points = rank_result if rank_result else (0, 0)
        courses.append(CourseProgressOut(course_id=course_id, course_title=titles.get(course_id, ""), progress=progress, rank=rank, points=points))
    courses.sort(key=lambda c: c.course_title)

    attendance = live_svc.attendance_summary_for_student(db, student.id)
    attendance_total = attendance.pop("total")

    exam_rows = db.execute(
        select(ExamAttempt.exam_id, func.max(ExamAttempt.percentage))
        .where(ExamAttempt.user_id == student.id, ExamAttempt.status.in_((AttemptStatus.submitted, AttemptStatus.expired)))
        .group_by(ExamAttempt.exam_id)
    ).all()
    exam_ids = [row[0] for row in exam_rows]
    exam_titles = dict(db.execute(select(Exam.id, Exam.title).where(Exam.id.in_(exam_ids))).all()) if exam_ids else {}
    latest_by_exam: dict[int, ExamAttempt] = {}
    attempts_used_by_exam: dict[int, int] = {}
    if exam_ids:
        all_attempts = db.scalars(
            select(ExamAttempt).where(ExamAttempt.exam_id.in_(exam_ids), ExamAttempt.user_id == student.id, ExamAttempt.status.in_((AttemptStatus.submitted, AttemptStatus.expired)))
            .order_by(ExamAttempt.submitted_at.desc())
        ).all()
        for a in all_attempts:
            attempts_used_by_exam[a.exam_id] = attempts_used_by_exam.get(a.exam_id, 0) + 1
            latest_by_exam.setdefault(a.exam_id, a)
    exams = [
        ExamSummaryOut(
            exam_id=exam_id, exam_title=exam_titles.get(exam_id, ""), best_percentage=best,
            latest_percentage=latest_by_exam[exam_id].percentage if exam_id in latest_by_exam else None,
            attempts_used=attempts_used_by_exam.get(exam_id, 0),
        )
        for exam_id, best in exam_rows
    ]

    history, _ = points_svc.point_history(db, student.id, page=1, page_size=10)

    whatsapp_lines = [
        f"الطالب: {student.full_name} ({student.student_code or ''})",
        f"التقدم العام: {round(sum(c.progress for c in courses) / len(courses)) if courses else 0}%",
        f"نسبة الحضور: {attendance['rate']}%",
        f"إجمالي النقاط: {profile.total_points or 0}",
    ]

    return StudentReportDetail(
        student=_report_student_out(student, profile),
        subscription=profile.subscription_status.value, subscription_expires_at=profile.subscription_expires_at,
        groups=groups, courses=courses,
        attendance=ReportAttendanceOut(present=attendance["present"], late=attendance["late"], absent=attendance["absent"], excused=attendance["excused"], rate=attendance["rate"]),
        attendance_records_count=attendance_total, exams=exams, points_total=profile.total_points or 0, points_recent=history,
        last_login_at=student.last_login_at, admin_notes=profile.admin_notes or "", whatsapp_summary_text="\n".join(whatsapp_lines),
    )


def export_student_report_detail_xlsx(detail: StudentReportDetail, *, academy_name: str = "Talent Academy") -> bytes:
    wb = Workbook()
    student_title = f"{detail.student.full_name} ({detail.student.student_code or '-'})"

    summary = wb.active
    summary.title = "Summary"
    write_report_sheet(
        summary, report_title=f"Student Report — {student_title}", headers=["Field", "Value"],
        rows=[
            ["Full Name", detail.student.full_name], ["Student ID", detail.student.student_code or ""],
            ["Grade", detail.student.grade_level or ""], ["Type", detail.student.student_type],
            ["Subscription", detail.subscription], ["Attendance %", detail.attendance.rate],
            ["Total Points", detail.points_total],
            ["Last Login", detail.last_login_at.strftime("%Y-%m-%d %H:%M") if detail.last_login_at else "Never"],
        ],
        academy_name=academy_name, column_widths=[24, 36],
    )

    courses_ws = wb.create_sheet("Courses")
    write_report_sheet(
        courses_ws, report_title=f"Courses — {student_title}", headers=["Course", "Progress %", "Rank", "Points"],
        rows=[[c.course_title, c.progress, c.rank, c.points] for c in detail.courses],
        academy_name=academy_name, column_widths=[30, 12, 8, 10],
    )

    exams_ws = wb.create_sheet("Exams")
    write_report_sheet(
        exams_ws, report_title=f"Exams — {student_title}", headers=["Exam", "Best %", "Latest %", "Attempts"],
        rows=[[e.exam_title, e.best_percentage, e.latest_percentage, e.attempts_used] for e in detail.exams],
        academy_name=academy_name, column_widths=[30, 10, 10, 10],
    )

    points_ws = wb.create_sheet("Points")
    write_report_sheet(
        points_ws, report_title=f"Points — {student_title}", headers=["Date", "Description", "Course", "Points"],
        rows=[[p.created_at.strftime("%Y-%m-%d %H:%M"), p.description, p.course_title or "", p.points] for p in detail.points_recent],
        academy_name=academy_name, totals_row=["", "", "Total", sum(p.points for p in detail.points_recent)] if detail.points_recent else None,
        column_widths=[16, 34, 22, 10],
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------- period report (PDF §5.4) ----

def student_period_report(db: Session, academy_id: int, student_id: int, date_from: datetime | None, date_to: datetime | None) -> StudentPeriodReport:
    student = get_student_or_404(db, academy_id, student_id)
    profile = student.student_profile

    course_ids = accessible_course_ids(db, student.id)
    titles = dict(db.execute(select(Course.id, Course.title).where(Course.id.in_(course_ids))).all()) if course_ids else {}
    courses = []
    for course_id in course_ids:
        progress = course_progress_percent(db, student.id, course_id)
        rank_result = points_svc.course_rank_for(db, course_id, student.id)
        rank, points = rank_result if rank_result else (0, 0)
        courses.append(CourseProgressOut(course_id=course_id, course_title=titles.get(course_id, ""), progress=progress, rank=rank, points=points))
    courses.sort(key=lambda c: c.course_title)

    attendance = live_svc.attendance_summary_for_student(db, student.id, since=date_from, until=date_to)
    attendance.pop("total", None)

    exam_stmt = select(ExamAttempt, Exam.title).join(Exam, Exam.id == ExamAttempt.exam_id).where(
        ExamAttempt.user_id == student.id, ExamAttempt.status.in_((AttemptStatus.submitted, AttemptStatus.expired))
    )
    if date_from is not None:
        exam_stmt = exam_stmt.where(ExamAttempt.submitted_at >= date_from)
    if date_to is not None:
        exam_stmt = exam_stmt.where(ExamAttempt.submitted_at <= date_to)
    exam_rows = db.execute(exam_stmt.order_by(ExamAttempt.submitted_at.desc())).all()
    exams = [PeriodExamOut(exam_title=title, percentage=a.percentage, passed=a.passed, submitted_at=a.submitted_at) for a, title in exam_rows]

    points_stmt = select(PointLedger).where(PointLedger.user_id == student.id)
    if date_from is not None:
        points_stmt = points_stmt.where(PointLedger.created_at >= date_from)
    if date_to is not None:
        points_stmt = points_stmt.where(PointLedger.created_at <= date_to)
    point_rows = db.scalars(points_stmt.order_by(PointLedger.created_at.desc())).all()
    point_course_ids = {r.course_id for r in point_rows if r.course_id}
    point_titles = dict(db.execute(select(Course.id, Course.title).where(Course.id.in_(point_course_ids))).all()) if point_course_ids else {}
    points_events = [
        PointEvent(
            id=r.id, event_key=r.event_key, source_type=r.source_type.value, source_id=r.source_id, description=r.description,
            points=r.points, course_id=r.course_id, course_title=point_titles.get(r.course_id), created_at=r.created_at,
        )
        for r in point_rows
    ]

    lesson_stmt = (
        select(LessonProgress.completed_at, Lesson.title, Course.title)
        .select_from(LessonProgress)
        .join(Lesson, Lesson.id == LessonProgress.lesson_id)
        .join(Section, Section.id == Lesson.section_id)
        .join(Course, Course.id == Section.course_id)
        .where(LessonProgress.user_id == student.id, LessonProgress.completed.is_(True))
    )
    if date_from is not None:
        lesson_stmt = lesson_stmt.where(LessonProgress.completed_at >= date_from)
    if date_to is not None:
        lesson_stmt = lesson_stmt.where(LessonProgress.completed_at <= date_to)
    lesson_rows = db.execute(lesson_stmt.order_by(LessonProgress.completed_at.desc())).all()
    completed_lessons = [PeriodLessonOut(lesson_title=lt, course_title=ct, completed_at=at) for at, lt, ct in lesson_rows if at is not None]

    return StudentPeriodReport(
        student=_report_student_out(student, profile), period=PeriodOut(date_from=date_from, date_to=date_to),
        courses=courses, attendance=ReportAttendanceOut(**attendance), exams=exams,
        points_total=sum(p.points for p in points_events), points_events=points_events,
        completed_lessons=completed_lessons, admin_notes=profile.admin_notes or "",
    )
