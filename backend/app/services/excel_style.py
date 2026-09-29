"""Phase 9: shared openpyxl styling so every export (students, attendance, exam results, PDF's
sibling xlsx exports) looks like one product instead of five bare `ws.append(...)` dumps -- title
block with academy branding, styled header row, frozen header, sane column widths, RTL sheet
direction, and an optional totals row.
"""

from __future__ import annotations

from datetime import datetime, timezone
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

PRIMARY_FILL = PatternFill("solid", fgColor="1A5AD9")
TOTALS_FILL = PatternFill("solid", fgColor="EEF2FA")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1A5AD9")
SUBTITLE_FONT = Font(size=10, color="666666")
TOTALS_FONT = Font(bold=True)
THIN = Side(style="thin", color="D9D9D9")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_report_sheet(
    ws: Worksheet,
    *,
    report_title: str,
    headers: list[str],
    rows: list[list[Any]],
    academy_name: str = "Talent Academy",
    date_range_text: str = "",
    generated_at: datetime | None = None,
    column_widths: list[int] | None = None,
    totals_row: list[Any] | None = None,
    rtl: bool = True,
) -> None:
    """Writes a fully-styled sheet in place: 3-row branding/title block, a styled+frozen header
    row, the data rows, and an optional bold totals row -- the shape every export in this app now
    shares (spec: "Include academy branding, report title, selected date range, generated-at
    date, styled headers ... frozen rows, correct column widths, totals, RTL where appropriate")."""
    ws.sheet_view.rightToLeft = rtl

    ws.cell(row=1, column=1, value=academy_name).font = TITLE_FONT
    ws.cell(row=2, column=1, value=report_title).font = Font(bold=True, size=12)
    meta_bits = []
    if date_range_text:
        meta_bits.append(date_range_text)
    meta_bits.append(f"Generated: {(generated_at or datetime.now(timezone.utc)).strftime('%Y-%m-%d %H:%M UTC')}")
    ws.cell(row=3, column=1, value=" | ".join(meta_bits)).font = SUBTITLE_FONT
    for r in (1, 2, 3):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(len(headers), 1))

    header_row = 5
    for col, label in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=label)
        cell.font = HEADER_FONT
        cell.fill = PRIMARY_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = CELL_BORDER
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.row_dimensions[header_row].height = 20

    data_row = header_row
    for row_values in rows:
        data_row += 1
        for col, value in enumerate(row_values, start=1):
            cell = ws.cell(row=data_row, column=col, value=value)
            cell.border = CELL_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")

    if totals_row is not None:
        data_row += 1
        for col, value in enumerate(totals_row, start=1):
            cell = ws.cell(row=data_row, column=col, value=value)
            cell.font = TOTALS_FONT
            cell.fill = TOTALS_FILL
            cell.border = CELL_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = column_widths or [max(12, min(32, len(h) + 4)) for h in headers]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{header_row}"
